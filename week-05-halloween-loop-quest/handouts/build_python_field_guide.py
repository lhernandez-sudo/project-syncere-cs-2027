from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.colors import HexColor, white, black
from reportlab.pdfbase.pdfmetrics import stringWidth

OUT = Path(__file__).parent
DOCX = OUT / "python-field-guide-editable.docx"
PDF = OUT / "python-field-guide-print-booklet.pdf"

PAGES = [
    ("Python Field Guide", "A beginner lookup guide", [
        ("Use this book while you code", "Look up a pattern, copy its shape, then change the names and values for your program."),
        ("You do not need to memorize everything", "Programmers look up syntax. Your job is to understand what the code is doing and test it."),
        ("Concepts inside", "Variables  Lists  Conditions  For loops  While loops  Functions  return  None"),
    ]),
    ("How to read Python", "Trace one line at a time", [
        ("Assignment", "coins = 3 means store 3 in the name coins. The = sign assigns a value."),
        ("Indentation", "Indented lines belong to the if, loop, or function above them."),
        ("Function call", "print(coins) runs the function named print with coins as its argument."),
        ("Predict then run", "Before you press Run, predict the output. Then compare your prediction with Python."),
    ]),
    ("Variables and lists", "Store one value or a collection", [
        ("Variables", "name = \"Luigi\"\nlives = 3\nis_ready = True"),
        ("Lists", "items = [\"key\", \"star\"]\nitems.append(\"mushroom\")"),
        ("List lookup", "items[0] is the first item. Python starts counting indexes at 0."),
        ("Useful checks", "len(items) counts items.  value in items checks whether a value is present."),
    ]),
    ("Conditions and logic", "Choose which code runs", [
        ("Basic pattern", "if coins >= 10:\n    print(\"1-UP\")\nelse:\n    print(\"Keep going\")"),
        ("Comparisons", "== equal   != not equal   > greater   < less   >= at least   <= at most"),
        ("Logic", "and needs both conditions.  or needs either condition.  not reverses True and False."),
        ("Remember", "Use == to compare. Use = to assign a value."),
    ]),
    ("For loops and range", "Repeat a known number of times", [
        ("Loop through a list", "for enemy in [\"Boo\", \"Goomba\"]:\n    print(enemy)"),
        ("What range means", "range(4) produces 0, 1, 2, 3. It starts at 0 and stops before 4."),
        ("Three useful forms", "range(4) -> 0,1,2,3\nrange(2, 5) -> 2,3,4\nrange(2, 9, 2) -> 2,4,6,8"),
        ("Accumulator", "total = 0\nfor coin in [2, 4, 6]:\n    total = total + coin"),
    ]),
    ("While loops", "Repeat while a condition stays True", [
        ("Basic pattern", "lives = 3\nwhile lives > 0:\n    lives = lives - 1"),
        ("Ask two questions", "What condition keeps the loop running? What changes so the loop can stop?"),
        ("Avoid infinite loops", "If the condition never becomes False, the loop never stops. Update the condition variable inside the loop."),
        ("Choose the loop", "Use for for a list or known count. Use while when repetition depends on a condition."),
    ]),
    ("Functions return and None", "Give reusable code a name", [
        ("Define and call", "def power(base, bonus):\n    return base + bonus\n\nscore = power(7, 5)"),
        ("Vocabulary", "Parameters are names in the definition. Arguments are values sent into the call."),
        ("return", "return sends a value back to the caller. Here, power(7, 5) returns 12."),
        ("None", "None means no value. A function without return still runs, but its result is None."),
        ("Print is different", "print shows text on screen. return gives a value back so the program can reuse it."),
    ]),
    ("Debug and look it up", "A routine for getting unstuck", [
        ("1 Read the error", "The last line usually names the error and points toward the problem."),
        ("2 Check the shape", "Look for missing colons, parentheses, quotes, indentation, or misspelled names."),
        ("3 Trace values", "Use print(variable) to see what a value is during the program."),
        ("4 Test a small piece", "Run one line, one condition, or one function call before testing the whole program."),
        ("Good search", "Search the exact goal: Python add item to list or Python range start stop step."),
        ("Final check", "Can you explain what each variable stores, when each loop stops, and what each function returns?"),
    ]),
]

# Sixteen-page edition. Current-course concepts come first; later pages are a
# preview/reference for future lessons rather than required prior knowledge.
PAGES = [
    ("Python Field Guide", "A black and white lookup guide for learning Python", [
        ("How to use this guide", "Find the idea you need, copy the pattern, change the names and values, then run a small test."),
        ("Looking things up is part of programming", "You do not need to memorize every command. You do need to explain what your code is doing."),
        ("Start with what we have learned", "Pages 2 through 8 review our current course. Pages 9 through 16 prepare you for future lessons."),
    ]),
    ("What we have learned", "Current course reference", [
        ("Values and variables", "Store text, numbers, and True or False values under useful names."),
        ("Program decisions", "Use if, elif, else, comparisons, and logic to choose which code runs."),
        ("Repetition", "Use for with a collection or range. Use while when repetition depends on a condition."),
        ("Reusable code", "Use functions, parameters, arguments, and return values to organize a program."),
        ("Core habit", "Predict the result, run the code, compare, and investigate any difference."),
    ]),
    ("Variables lists and input", "Store information and receive user choices", [
        ("Variables", "player = \"Mario\"\nlives = 3\nis_ready = True"),
        ("Input returns text", "name = input(\"Your name: \")\nprint(\"Hello\", name)"),
        ("Lists", "items = [\"key\", \"star\"]\nitems.append(\"mushroom\")\nfirst = items[0]"),
        ("Useful list tools", "len(items) counts items. value in items checks membership. items.remove(value) removes a matching item."),
        ("Index reminder", "The first list index is 0. The last item can be read with items[-1]."),
    ]),
    ("Conditions and logic", "Choose which code runs", [
        ("Basic pattern", "if coins >= 10:\n    print(\"1-UP\")\nelif coins >= 5:\n    print(\"Almost\")\nelse:\n    print(\"Keep going\")"),
        ("Comparisons", "== equal   != not equal   > greater   < less   >= at least   <= at most"),
        ("Logic", "and needs both conditions. or needs either condition. not reverses True and False."),
        ("Common mistake", "Use == to compare values. Use = to assign a value to a variable."),
    ]),
    ("For loops and range", "Repeat a known number of times", [
        ("Loop through a list", "for enemy in [\"Boo\", \"Goomba\"]:\n    print(enemy)"),
        ("What range means", "range(4) produces 0, 1, 2, 3. It starts at 0 and stops before 4."),
        ("Three useful forms", "range(4) -> 0,1,2,3\nrange(2, 5) -> 2,3,4\nrange(2, 9, 2) -> 2,4,6,8"),
        ("Accumulator pattern", "total = 0\nfor coin in [2, 4, 6]:\n    total = total + coin"),
    ]),
    ("While loops", "Repeat while a condition stays True", [
        ("Basic pattern", "lives = 3\nwhile lives > 0:\n    lives = lives - 1"),
        ("Ask two questions", "What condition keeps the loop running? What changes so that condition becomes False?"),
        ("Avoid infinite loops", "If the condition never becomes False, the loop never stops. Update the condition variable inside the loop."),
        ("Choose the loop", "Use for for a list or known count. Use while when repetition depends on a condition."),
        ("Useful controls", "break exits a loop early. continue skips to the next iteration."),
    ]),
    ("Functions return and None", "Give reusable code a clear name", [
        ("Define and call", "def power(base, bonus):\n    return base + bonus\n\nscore = power(7, 5)"),
        ("Vocabulary", "Parameters are names in the definition. Arguments are values sent into the call."),
        ("return", "return sends a value back to the caller. Here, power(7, 5) returns 12."),
        ("None", "None means no value. A function without return still runs, but its result is None."),
        ("Print is different", "print displays information. return gives a value back so the program can reuse it."),
    ]),
    ("Current course quick check", "Use this before asking for help", [
        ("Trace", "Write the value of each variable after every line. For a loop, trace one iteration at a time."),
        ("Check the shape", "Look for missing colons, parentheses, quotes, indentation, or misspelled names."),
        ("Test one piece", "Run one expression, condition, loop, or function call before testing the whole program."),
        ("Explain", "Can you say what each variable stores, when each loop stops, and what each function returns?"),
        ("Search precisely", "Try: Python add item to list, Python range start stop step, or Python function returns None."),
    ]),
    ("Future strings and numbers", "Convert and format information", [
        ("Input conversion", "age_text = input(\"Age: \")\nage = int(age_text)\nheight = float(input(\"Height: \"))"),
        ("String methods", "name.strip() removes edge spaces. name.lower() changes case. text.split(\",\") makes a list."),
        ("Formatted strings", "coins = 12\nmessage = f\"You have {coins} coins\""),
        ("Type check", "type(value) reports a value's type. Common types are str, int, float, bool, list, and NoneType."),
    ]),
    ("Future dictionaries and sets", "Store labeled data and unique values", [
        ("Dictionary", "player = {\"name\": \"Luigi\", \"lives\": 3}\nprint(player[\"name\"])"),
        ("Safe lookup", "player.get(\"power\", \"small\") returns a default if the key is missing."),
        ("Loop through pairs", "for key, value in player.items():\n    print(key, value)"),
        ("Set", "unique_items = set(items) removes duplicates. Sets are useful for membership checks."),
    ]),
    ("Future function patterns", "Make code easier to reuse and test", [
        ("Default argument", "def greet(name, greeting=\"Hello\"):\n    return f\"{greeting}, {name}\""),
        ("Small focused functions", "A function should usually do one clear job. A good name describes that job."),
        ("Local variables", "A variable created inside a function is normally local to that function."),
        ("Documentation", "Place a short docstring under def to explain the inputs, behavior, and returned value."),
    ]),
    ("Future errors and exceptions", "Handle expected problems", [
        ("Read the traceback", "Start at the final line for the error type and message, then check the named line in your code."),
        ("Common errors", "SyntaxError: code shape is invalid. NameError: name is unknown. TypeError: operation does not fit the value type."),
        ("Expected input problem", "try:\n    age = int(input(\"Age: \"))\nexcept ValueError:\n    print(\"Enter a whole number\")"),
        ("Do not hide every error", "Catch errors you expect and can handle. Unexpected errors should stay visible while debugging."),
    ]),
    ("Future files and modules", "Save data and use organized code", [
        ("Read a text file", "with open(\"notes.txt\", \"r\") as file:\n    text = file.read()"),
        ("Write a text file", "with open(\"score.txt\", \"w\") as file:\n    file.write(str(score))"),
        ("Import a module", "import random\nchoice = random.choice(items)"),
        ("Main guard", "if __name__ == \"__main__\":\n    main()"),
        ("File habit", "Use with so Python closes the file even if something goes wrong."),
    ]),
    ("Future classes and objects", "Bundle related data and behavior", [
        ("Class pattern", "class Player:\n    def __init__(self, name):\n        self.name = name\n        self.lives = 3"),
        ("Create an object", "hero = Player(\"Mario\")\nprint(hero.name)"),
        ("Method", "def lose_life(self):\n    self.lives -= 1"),
        ("Mental model", "A class is a blueprint. An object is one value created from that blueprint. self means this object."),
    ]),
    ("Future testing and design", "Build reliable programs in small steps", [
        ("Arrange act assert", "Arrange inputs, call the function, then check the result you expected."),
        ("Simple self check", "assert power(7, 5) == 12\nassert double(0) == 0"),
        ("Test edges", "Try empty lists, zero, negative numbers, missing values, and the smallest or largest allowed input."),
        ("Decompose", "Break a large problem into named functions. Test each function before joining them together."),
        ("Comments", "Explain why a decision exists. Let clear names and small functions explain what the code does."),
    ]),
    ("Python lookup index", "Where to turn when you get stuck", [
        ("Current course", "Variables and input 3   Conditions 4   For and range 5   While 6   Functions and None 7   Quick check 8"),
        ("Future reference", "Strings 9   Dictionaries 10   Function patterns 11   Errors 12   Files 13   Classes 14   Testing 15"),
        ("Search formula", "Search Python plus your goal plus the value type. Example: Python remove item from list."),
        ("Ask a strong question", "Share what you expected, what happened, the smallest code that reproduces it, and the full error message."),
        ("Keep learning", "Read examples actively: predict, run, change one thing, and explain the new result."),
    ]),
]

ACCENT = "000000"
INK = "000000"
MUTED = "444444"
PALE = "F5F5F5"

CODE_STARTS = (
    "def ", "class ", "if ", "elif ", "else:", "for ", "while ",
    "try:", "except ", "with ", "import ", "from ", "return ",
    "print(", "assert ", "break", "continue",
)

def is_code_block(body):
    """Only style actual multi-line snippets as code, not prose mentioning code."""
    lines = body.split("\n")
    if len(lines) < 2:
        return False
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith(CODE_STARTS) or " = " in stripped or ".append(" in stripped:
            return True
    return False

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def add_run(p, text, bold=False, size=9, color=INK, font="Aptos"):
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = font
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    return r

def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(5.5), Inches(8.5)
    sec.top_margin = sec.bottom_margin = Inches(.48)
    sec.left_margin = sec.right_margin = Inches(.48)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(9)
    for page_no, (title, subtitle, cards) in enumerate(PAGES, 1):
        if page_no > 1:
            doc.add_page_break()
        p = doc.add_paragraph(); p.style = styles["Title"]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, title, True, 20, "000000")
        p.paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph(); add_run(p, subtitle, False, 10, MUTED)
        p.paragraph_format.space_after = Pt(9)
        for heading, body in cards:
            t = doc.add_table(rows=1, cols=1); t.autofit = False; t.columns[0].width = Inches(4.5)
            cell = t.cell(0,0); set_cell_shading(cell, PALE)
            cell.margin_top = cell.margin_bottom = 90; cell.margin_left = cell.margin_right = 110
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
            add_run(p, heading, True, 10, ACCENT)
            lines = body.split("\n")
            for i, line in enumerate(lines):
                p = cell.add_paragraph() if i or p.text else p
                p.paragraph_format.space_after = Pt(1)
                is_code = is_code_block(body)
                add_run(p, line, False, 8.4 if is_code else 8.8, INK, "Consolas" if is_code else "Aptos")
            spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(3)
        footer = doc.add_paragraph(); footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_run(footer, f"PYTHON FIELD GUIDE   {page_no}", True, 7.5, MUTED)
    doc.core_properties.title = "Python Field Guide"
    doc.core_properties.subject = "Beginner Python syntax and concept lookup guide"
    doc.save(DOCX)

def wrap(canvas, text, font, size, width):
    words = text.split(); lines=[]; line=""
    for word in words:
        trial = (line + " " + word).strip()
        if stringWidth(trial, font, size) <= width: line=trial
        else: lines.append(line); line=word
    if line: lines.append(line)
    return lines

def wrap_code(text, font, size, width):
    """Wrap code without destroying indentation, which is Python syntax."""
    text = text.expandtabs(4)
    if stringWidth(text, font, size) <= width:
        return [text]
    indent = text[:len(text) - len(text.lstrip(" "))]
    continuation = indent + "    "
    remaining = text
    lines = []
    first = True
    while remaining:
        prefix = "" if first else continuation
        room = max(8, int((width - stringWidth(prefix, font, size)) / stringWidth("M", font, size)))
        cut = min(len(remaining), room)
        if cut < len(remaining):
            natural = max(remaining.rfind(" ", 0, cut), remaining.rfind(",", 0, cut) + 1)
            if natural > max(8, cut // 2):
                cut = natural
        lines.append(prefix + remaining[:cut].rstrip())
        remaining = remaining[cut:].lstrip()
        first = False
    return lines

def draw_logical_page(c, x0, y0, title, subtitle, cards, page_no):
    w, h = 5.5*72, 8.5*72
    c.setFillColor(white); c.rect(x0,y0,w,h,fill=1,stroke=0)
    c.setFillColor(HexColor("#000000"))
    title_size = 19
    while stringWidth(title, "Helvetica-Bold", title_size) > w - 76 and title_size > 11:
        title_size -= .5
    c.setFont("Helvetica-Bold",title_size); c.drawString(x0+34,y0+h-48,title)
    c.setFillColor(HexColor("#556070")); c.setFont("Helvetica",9.5); c.drawString(x0+34,y0+h-65,subtitle)
    y=y0+h-88
    available = y-(y0+36); gap=7
    heights=[]
    for heading, body in cards:
        code = is_code_block(body)
        font="Courier" if code else "Helvetica"; size=7.5 if code else 8
        n=sum(max(1,len(wrap(c,line,font,size,w-82))) for line in body.split("\n"))
        heights.append(28+n*10)
    scale=min(1,available/(sum(heights)+gap*(len(cards)-1)))
    for (heading,body),raw_h in zip(cards,heights):
        ch=raw_h*scale
        c.setFillColor(white); c.setStrokeColor(HexColor("#B8B8B8")); c.setLineWidth(.7)
        c.roundRect(x0+34,y-ch,w-68,ch,5,fill=1,stroke=1)
        c.setFillColor(black); c.setFont("Helvetica-Bold",9); c.drawString(x0+43,y-15,heading)
        code = is_code_block(body)
        font="Courier" if code else "Helvetica"; size=7.5 if code else 8
        ty=y-28; c.setFont(font,size); c.setFillColor(black)
        for source_line in body.split("\n"):
            rendered_lines = wrap_code(source_line,font,size,w-86) if code else wrap(c,source_line,font,size,w-86)
            for line in rendered_lines or [""]:
                c.drawString(x0+43,ty,line); ty-=10*scale
        y-=ch+gap
    c.setFillColor(HexColor("#556070")); c.setFont("Helvetica-Bold",7)
    c.drawRightString(x0+w-34,y0+20,f"PYTHON FIELD GUIDE   {page_no}")

def build_booklet_pdf():
    c=Canvas(str(PDF),pagesize=landscape(letter))
    # Booklet imposition for any page count divisible by four. The face order
    # is ready for duplex printing on the short edge, folding, then stapling.
    count = len(PAGES)
    assert count % 4 == 0
    faces=[]
    for sheet in range(count // 4):
        faces.append((count - sheet * 2, 1 + sheet * 2))
        faces.append((2 + sheet * 2, count - 1 - sheet * 2))
    for left,right in faces:
        draw_logical_page(c,0,0,*PAGES[left-1],left)
        draw_logical_page(c,5.5*72,0,*PAGES[right-1],right)
        c.setStrokeColor(HexColor("#C8CDD5")); c.setDash(2,3); c.line(5.5*72,0,5.5*72,8.5*72); c.setDash()
        c.showPage()
    c.save()

if __name__ == "__main__":
    OUT.mkdir(parents=True,exist_ok=True)
    build_docx(); build_booklet_pdf()
    print(DOCX); print(PDF)
